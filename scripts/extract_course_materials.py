"""Convert course source documents and Swift examples into RAG-ready Markdown.

Usage:
  python scripts/extract_course_materials.py SOURCE_DIRECTORY knowledge

The source directory is intentionally kept outside the repository. Generated
Markdown contains source filenames and page numbers, but no absolute paths.
"""

from __future__ import annotations

import hashlib
import re
import sys
from pathlib import Path

import pdfplumber
from docx import Document


def slugify(value: str) -> str:
    value = value.lower().replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value or "source"


def clean_text(value: str) -> str:
    value = value.replace("\r", "").replace("\u00a0", " ")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def scrub_professor_information(value: str) -> str:
    value = re.sub(r"^[A-Z][a-z]+ [A-Z][a-z]+ IT Help$", "IT Help", value, flags=re.MULTILINE)
    value = re.sub(r"^Office Hours:.*?Hours:\s*(.*?)$", r"Hours: \1", value, flags=re.MULTILINE)
    value = re.sub(
        r"^[^\n]*@usc\.edu Contact Info:\s*([^\n]+)$",
        r"Contact Info: \1",
        value,
        flags=re.MULTILINE,
    )
    value = re.sub(r"^//\s*Created by .*? on .*?\n", "", value, flags=re.MULTILINE)
    return value


def extract_pdf(source: Path, output_directory: Path) -> tuple[Path, int, int]:
    is_course_syllabus = source.name.lower().startswith("idsn533 ios app dev")
    title = "IDSN 533 iOS Application Design and Development Fall 2026" if is_course_syllabus else source.stem
    source_label = "Course syllabus PDF" if is_course_syllabus else source.name
    sections = [f"# {title}", "", f"Source document: `{source_label}`"]
    extracted_characters = 0

    with pdfplumber.open(source) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = scrub_professor_information(clean_text(page.extract_text(x_tolerance=2, y_tolerance=3) or ""))
            sections.extend(["", f"## Page {page_number}", ""])
            if text:
                sections.append(text)
                extracted_characters += len(text)
            else:
                sections.append("[No machine-readable text was found on this page.]")
        page_count = len(pdf.pages)

    destination_name = "course-syllabus-fall-2026.md" if is_course_syllabus else f"{slugify(source.stem)}.md"
    destination = output_directory / destination_name
    destination.write_text("\n".join(sections).strip() + "\n", encoding="utf-8")
    return destination, page_count, extracted_characters


def markdown_table(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]

    def escape(cell: str) -> str:
        return clean_text(cell).replace("\n", "<br>").replace("|", "\\|")

    output = ["| " + " | ".join(escape(cell) for cell in normalized[0]) + " |"]
    output.append("| " + " | ".join(["---"] * width) + " |")
    output.extend("| " + " | ".join(escape(cell) for cell in row) + " |" for row in normalized[1:])
    return output


def extract_docx(source: Path, output_directory: Path) -> tuple[Path, int, int]:
    document = Document(source)
    sections = [f"# {source.stem}", "", f"Source document: `{source.name}`", ""]
    paragraph_count = 0

    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        paragraph_count += 1
        style = (paragraph.style.name or "").lower()
        inferred_heading = re.match(
            r"^(Exercise\s+\d+|Midterm Project|Research Project|Final Project|Final Pitch|Final Prototype|Final Beta|Final Appstore Submission|Final Presentation|Final Documentation|Participation|Extra Credit|Grading Breakdown)\b",
            text,
            flags=re.IGNORECASE,
        )
        if style.startswith("heading"):
            match = re.search(r"(\d+)", style)
            level = min(6, (int(match.group(1)) if match else 2) + 1)
            sections.extend([f"{'#' * level} {text}", ""])
        elif style == "title":
            sections.extend([f"## {text}", ""])
        elif inferred_heading:
            sections.extend([f"## {text}", ""])
        elif style.startswith("list"):
            sections.append(f"- {text}")
        else:
            sections.extend([text, ""])

    for table_number, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        sections.extend([f"## Table {table_number}", "", *markdown_table(rows), ""])

    destination = output_directory / f"{slugify(source.stem)}.md"
    content = "\n".join(sections).strip() + "\n"
    destination.write_text(content, encoding="utf-8")
    return destination, paragraph_count, len(content)


def extract_swift_examples(source_directory: Path, output_directory: Path) -> list[tuple[Path, int, int]]:
    candidates = sorted(source_directory.rglob("*.swift"))
    if not candidates:
        return []

    repository_root = source_directory
    nested_root = source_directory / "swift-examples-main"
    if nested_root.is_dir():
        repository_root = nested_root

    seen_hashes: set[str] = set()
    topics: dict[str, list[tuple[str, str]]] = {}
    duplicate_count = 0

    for source in candidates:
        content = scrub_professor_information(clean_text(source.read_text(encoding="utf-8", errors="replace")))
        if not content:
            continue
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
        if digest in seen_hashes:
            duplicate_count += 1
            continue
        seen_hashes.add(digest)

        relative = source.relative_to(repository_root)
        topic = relative.parts[0] if len(relative.parts) > 1 else "general"
        topics.setdefault(topic, []).append((relative.as_posix(), content))

    results = []
    for topic, examples in sorted(topics.items()):
        sections = [
            f"# Swift Examples: {topic.replace('-', ' ').title()}",
            "",
            "Source collection: `swift-examples-main`",
            "",
            f"This file contains {len(examples)} unique Swift source examples. Exact duplicate files were removed during extraction.",
        ]
        for relative_path, content in examples:
            display_name = Path(relative_path).stem.replace("_", " ")
            sections.extend([
                "",
                f"## {display_name}",
                "",
                f"Original path: `{relative_path}`",
                "",
                "```swift",
                content,
                "```",
            ])

        destination = output_directory / f"swift-examples-{slugify(topic)}.md"
        rendered = "\n".join(sections).strip() + "\n"
        destination.write_text(rendered, encoding="utf-8")
        results.append((destination, len(examples), len(rendered)))

    print(f"Swift extraction: {len(seen_hashes)} unique files, {duplicate_count} exact duplicates skipped")
    return results


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: extract_course_materials.py SOURCE_DIRECTORY OUTPUT_DIRECTORY", file=sys.stderr)
        return 2

    source_directory = Path(sys.argv[1]).resolve()
    output_directory = Path(sys.argv[2]).resolve()
    if not source_directory.is_dir():
        print(f"Source directory does not exist: {source_directory}", file=sys.stderr)
        return 1
    output_directory.mkdir(parents=True, exist_ok=True)

    outputs: list[tuple[Path, int, int]] = []
    for source in sorted(source_directory.glob("*.pdf")):
        outputs.append(extract_pdf(source, output_directory))
    for source in sorted(source_directory.glob("*.docx")):
        outputs.append(extract_docx(source, output_directory))
    outputs.extend(extract_swift_examples(source_directory / "swift-examples-main", output_directory))

    print("Generated knowledge files:")
    for destination, units, characters in outputs:
        print(f"- {destination.name}: {units} source units, {characters:,} characters")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
